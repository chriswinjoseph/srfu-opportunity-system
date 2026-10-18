"""
Club files (CSV or Excel) are no longer written directly by this command.

The older importer used its own duplicate rule (name + suburb + state) and
bypassed the shared review workflow, so it is disabled. Clubs are collected
through the Admin Scraper page, or imported through the Admin Data Import page
using a workbook that carries an organisation_type of Club.

    python manage.py import_clubs path/to/clubs.xlsx --dry-run     (still reports what the file holds)
"""

import csv
from pathlib import Path

from django.core.management.base import BaseCommand, CommandError
from openpyxl import load_workbook


class Command(BaseCommand):
    help = (
        "Club import is disabled because it bypassed the shared duplicate and "
        "review rules. --dry-run only reports how many rows a file holds."
    )

    def add_arguments(self, parser):
        parser.add_argument("file_path", type=str, help="Path to a CSV or XLSX club dataset.")
        parser.add_argument("--dry-run", action="store_true", help="Count the rows in the file; nothing is saved.")

    def handle(self, *args, **options):
        path = Path(options["file_path"])

        if not path.exists():
            raise CommandError(f"File not found: {path}")

        if not options["dry_run"]:
            from intake import legacy

            raise legacy.direct_write_disabled(
                "import_clubs",
                "Use the Admin Scraper page for the approved public club directory, or the Admin Data Import "
                "page with an organisation_type column of Club. Use --dry-run to count a file's rows.",
            )

        if path.suffix.lower() == ".csv":
            with path.open(newline="", encoding="utf-8-sig") as handle:
                rows = sum(1 for _ in csv.DictReader(handle))
        elif path.suffix.lower() == ".xlsx":
            workbook = load_workbook(path, read_only=True, data_only=True)
            rows = max(workbook.active.max_row - 1, 0)
            workbook.close()
        else:
            raise CommandError("Only .csv and .xlsx files can be read.")

        self.stdout.write(f"{rows} row(s) found. Nothing was written to the database.")
