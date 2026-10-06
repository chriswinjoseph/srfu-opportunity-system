"""
Import Shane's starter dataset from the cleaned Excel workbook.

Usage:
    python manage.py import_seed path/to/file.xlsx

Dry run:
    python manage.py import_seed path/to/file.xlsx --dry-run

Expected workbook sheets:
    - Organisations
    - Contacts

The importer:
    - imports Bank records
    - imports Contact records
    - prevents duplicate banks
    - prevents duplicate contacts
    - maps Contacted / Partially Contacted -> contacted
    - maps Not Yet Contacted -> not_yet_contacted
    - preserves contacted status for existing banks
    - creates a basic Opportunity for contacted organisations
"""

from datetime import date, datetime

from django.contrib.contenttypes.models import ContentType
from django.core.exceptions import ValidationError
from django.core.management.base import BaseCommand, CommandError
from django.core.validators import validate_email

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

from outreach.models import Bank, Contact, Opportunity


CONTACTED_VALUES = {
    "contacted",
    "partially contacted",
}


def clean_text(value):
    if value is None:
        return ""

    return str(value).strip()


def clean_contact_name(value):
    value = clean_text(value)

    if value.lower() == "(unspecified)":
        return ""

    return value


def first_email(value):
    """
    Organisation rows may contain multiple emails separated by semicolons.
    Use the first valid email as the Bank public_email.
    """

    value = clean_text(value)

    if not value:
        return ""

    emails = [
        email.strip()
        for email in value.split(";")
        if email.strip()
    ]

    for email in emails:
        try:
            validate_email(email)
            return email
        except ValidationError:
            continue

    return ""


def normalise_status(value):
    value = clean_text(value).lower()

    if value in CONTACTED_VALUES:
        return "contacted"

    return "not_yet_contacted"


def excel_date(value, workbook_epoch):
    """
    Convert Excel dates or serial values to Python date objects.
    """

    if value in (None, ""):
        return None

    if isinstance(value, datetime):
        return value.date()

    if isinstance(value, date):
        return value

    if isinstance(value, (int, float)):
        try:
            converted = from_excel(
                value,
                epoch=workbook_epoch,
            )

            if isinstance(converted, datetime):
                return converted.date()

            if isinstance(converted, date):
                return converted

        except (ValueError, TypeError):
            return None

    return None


def worksheet_rows_as_dicts(sheet):
    """
    Convert an Excel worksheet into dictionaries using row 1 as headers.
    """

    rows = sheet.iter_rows(values_only=True)

    try:
        headers = next(rows)
    except StopIteration:
        return

    headers = [
        clean_text(header).lower().replace(" ", "_")
        for header in headers
    ]

    for row_number, values in enumerate(rows, start=2):
        row = dict(zip(headers, values))
        yield row_number, row


class Command(BaseCommand):
    help = "Import Shane's cleaned bank/contact dataset from Excel."

    def add_arguments(self, parser):
        parser.add_argument(
            "xlsx_path",
            type=str,
            help="Path to the cleaned .xlsx dataset.",
        )

        parser.add_argument(
            "--dry-run",
            action="store_true",
            help="Validate and report without writing to the database.",
        )

    def handle(self, *args, **options):
        xlsx_path = options["xlsx_path"]
        dry_run = options["dry_run"]

        try:
            workbook = load_workbook(
                xlsx_path,
                data_only=True,
            )
        except FileNotFoundError:
            raise CommandError(
                f"File not found: {xlsx_path}"
            )
        except Exception as exc:
            raise CommandError(
                f"Could not open workbook: {exc}"
            )

        required_sheets = {
            "Organisations",
            "Contacts",
        }

        missing_sheets = (
            required_sheets
            - set(workbook.sheetnames)
        )

        if missing_sheets:
            raise CommandError(
                "Workbook is missing required sheet(s): "
                + ", ".join(sorted(missing_sheets))
            )

        organisation_sheet = workbook["Organisations"]
        contact_sheet = workbook["Contacts"]

        bank_created = 0
        bank_duplicates = 0
        bank_errors = 0

        contact_created = 0
        contact_duplicates = 0
        contact_errors = 0

        opportunity_created = 0

        errors = []

        # Maps dataset organisation_id -> Bank object
        organisation_map = {}

        # -----------------------------------------------------
        # Import Organisations / Banks
        # -----------------------------------------------------

        for row_number, row in worksheet_rows_as_dicts(
            organisation_sheet
        ):
            dataset_org_id = clean_text(
                row.get("organisation_id")
            )

            bank_name = clean_text(
                row.get("bank_name")
            )

            state = clean_text(
                row.get("state")
            )

            overall_status = clean_text(
                row.get("overall_status")
            )

            if not bank_name:
                bank_errors += 1

                errors.append(
                    f"Organisations row {row_number}: "
                    "bank_name is missing."
                )

                continue

            if not state:
                bank_errors += 1

                errors.append(
                    f"Organisations row {row_number}: "
                    f"state is missing for '{bank_name}'."
                )

                continue

            contact_status = normalise_status(
                overall_status
            )

            public_email = first_email(
                row.get("emails")
            )

            existing_bank = (
                Bank.objects
                .filter(
                    bank_name__iexact=bank_name
                )
                .first()
            )

            # -------------------------------------------------
            # Existing Bank
            # -------------------------------------------------

            if existing_bank:
                bank_duplicates += 1

                if dataset_org_id:
                    organisation_map[
                        dataset_org_id
                    ] = existing_bank

                # Preserve previous outreach from Shane's data.
                if (
                    contact_status == "contacted"
                    and existing_bank.contact_status != "contacted"
                ):
                    if dry_run:
                        self.stdout.write(
                            self.style.WARNING(
                                f"Organisations row {row_number}: "
                                f"'{bank_name}' already exists — "
                                "status would be updated to Contacted."
                            )
                        )

                    else:
                        existing_bank.contact_status = "contacted"

                        existing_bank.save(
                            update_fields=[
                                "contact_status"
                            ]
                        )

                        self.stdout.write(
                            self.style.WARNING(
                                f"Organisations row {row_number}: "
                                f"'{bank_name}' already exists — "
                                "status updated to Contacted."
                            )
                        )

                # Make sure contacted organisations have an Opportunity.
                if contact_status == "contacted":
                    content_type = (
                        ContentType.objects
                        .get_for_model(Bank)
                    )

                    if dry_run:
                        opportunity_exists = (
                            Opportunity.objects.filter(
                                content_type=content_type,
                                object_id=existing_bank.pk,
                            ).exists()
                        )

                        if not opportunity_exists:
                            opportunity_created += 1

                    else:
                        _, created = (
                            Opportunity.objects.get_or_create(
                                content_type=content_type,
                                object_id=existing_bank.pk,
                                defaults={
                                    "status": "contacted",
                                    "outreach_method": "email",
                                },
                            )
                        )

                        if created:
                            opportunity_created += 1

                self.stdout.write(
                    self.style.WARNING(
                        f"Organisations row {row_number}: "
                        f"'{bank_name}' already exists — reused."
                    )
                )

                continue

            # -------------------------------------------------
            # New Bank
            # -------------------------------------------------

            if dry_run:
                bank_created += 1

                self.stdout.write(
                    f"Organisations row {row_number}: "
                    f"would create '{bank_name}' "
                    f"(state={state}, "
                    f"status={contact_status})"
                )

                continue

            bank = Bank.objects.create(
                bank_name=bank_name,

                # Current Bank model requires region.
                # Dataset has state only, so state is used as fallback.
                state=state,
                region=state,

                public_email=public_email,
                public_phone="",
                website_url="",
                source_url="",

                contact_status=contact_status,
                record_source="Shane Starter Dataset",
            )

            bank_created += 1

            if dataset_org_id:
                organisation_map[
                    dataset_org_id
                ] = bank

            if contact_status == "contacted":
                content_type = (
                    ContentType.objects
                    .get_for_model(Bank)
                )

                _, created = (
                    Opportunity.objects.get_or_create(
                        content_type=content_type,
                        object_id=bank.pk,
                        defaults={
                            "status": "contacted",
                            "outreach_method": "email",
                        },
                    )
                )

                if created:
                    opportunity_created += 1

        # -----------------------------------------------------
        # Dry-run Contacts
        # -----------------------------------------------------

        if dry_run:
            for row_number, row in worksheet_rows_as_dicts(
                contact_sheet
            ):
                organisation_id = clean_text(
                    row.get("organisation_id")
                )

                email = clean_text(
                    row.get("email")
                )

                if not organisation_id:
                    contact_errors += 1

                    errors.append(
                        f"Contacts row {row_number}: "
                        "organisation_id is missing."
                    )

                    continue

                if not email:
                    contact_errors += 1

                    errors.append(
                        f"Contacts row {row_number}: "
                        "email is missing."
                    )

                    continue

                try:
                    validate_email(email)

                except ValidationError:
                    contact_errors += 1

                    errors.append(
                        f"Contacts row {row_number}: "
                        f"invalid email '{email}'."
                    )

                    continue

                bank = organisation_map.get(
                    organisation_id
                )

                if bank is not None:
                    bank_content_type = (
                        ContentType.objects
                        .get_for_model(Bank)
                    )

                    duplicate_contact = (
                        Contact.objects.filter(
                            content_type=bank_content_type,
                            object_id=bank.pk,
                            email__iexact=email,
                        ).exists()
                    )

                    if duplicate_contact:
                        contact_duplicates += 1
                        continue

                contact_created += 1

            self._print_summary(
                dry_run=True,
                bank_created=bank_created,
                bank_duplicates=bank_duplicates,
                bank_errors=bank_errors,
                contact_created=contact_created,
                contact_duplicates=contact_duplicates,
                contact_errors=contact_errors,
                opportunity_created=opportunity_created,
                errors=errors,
            )

            return

        # -----------------------------------------------------
        # Import Contacts
        # -----------------------------------------------------

        bank_content_type = (
            ContentType.objects
            .get_for_model(Bank)
        )

        for row_number, row in worksheet_rows_as_dicts(
            contact_sheet
        ):
            dataset_org_id = clean_text(
                row.get("organisation_id")
            )

            bank_name = clean_text(
                row.get("bank_name")
            )

            contact_name = clean_contact_name(
                row.get("contact_name")
            )

            email = clean_text(
                row.get("email")
            )

            if not dataset_org_id:
                contact_errors += 1

                errors.append(
                    f"Contacts row {row_number}: "
                    "organisation_id is missing."
                )

                continue

            bank = organisation_map.get(
                dataset_org_id
            )

            # Fallback if bank existed before import
            # and no dataset mapping was available.
            if bank is None and bank_name:
                bank = (
                    Bank.objects
                    .filter(
                        bank_name__iexact=bank_name
                    )
                    .first()
                )

            if bank is None:
                contact_errors += 1

                errors.append(
                    f"Contacts row {row_number}: "
                    f"matching bank was not found "
                    f"for '{bank_name}'."
                )

                continue

            if not email:
                contact_errors += 1

                errors.append(
                    f"Contacts row {row_number}: "
                    f"email is missing for '{bank_name}'."
                )

                continue

            try:
                validate_email(email)

            except ValidationError:
                contact_errors += 1

                errors.append(
                    f"Contacts row {row_number}: "
                    f"invalid email '{email}'."
                )

                continue

            duplicate_contact = (
                Contact.objects.filter(
                    content_type=bank_content_type,
                    object_id=bank.pk,
                    email__iexact=email,
                ).exists()
            )

            if duplicate_contact:
                contact_duplicates += 1

                self.stdout.write(
                    self.style.WARNING(
                        f"Contacts row {row_number}: "
                        f"'{email}' already exists for "
                        f"'{bank.bank_name}' — skipped."
                    )
                )

                continue

            last_verified = excel_date(
                row.get("last_observed_date"),
                workbook.epoch,
            )

            Contact.objects.create(
                content_type=bank_content_type,
                object_id=bank.pk,
                contact_name=contact_name,
                role="",
                email=email,
                phone="",
                source_url="",
                last_verified=last_verified,
            )

            contact_created += 1

        self._print_summary(
            dry_run=False,
            bank_created=bank_created,
            bank_duplicates=bank_duplicates,
            bank_errors=bank_errors,
            contact_created=contact_created,
            contact_duplicates=contact_duplicates,
            contact_errors=contact_errors,
            opportunity_created=opportunity_created,
            errors=errors,
        )

    def _print_summary(
        self,
        *,
        dry_run,
        bank_created,
        bank_duplicates,
        bank_errors,
        contact_created,
        contact_duplicates,
        contact_errors,
        opportunity_created,
        errors,
    ):
        self.stdout.write("")

        self.stdout.write(
            self.style.SUCCESS(
                f"Banks created: {bank_created}"
            )
        )

        self.stdout.write(
            self.style.WARNING(
                f"Bank duplicates reused: "
                f"{bank_duplicates}"
            )
        )

        self.stdout.write(
            self.style.SUCCESS(
                f"Contacts created: {contact_created}"
            )
        )

        self.stdout.write(
            self.style.WARNING(
                f"Contact duplicates skipped: "
                f"{contact_duplicates}"
            )
        )

        self.stdout.write(
            self.style.SUCCESS(
                f"Opportunities created: "
                f"{opportunity_created}"
            )
        )

        if bank_errors:
            self.stdout.write(
                self.style.ERROR(
                    f"Bank rows with errors: "
                    f"{bank_errors}"
                )
            )

        if contact_errors:
            self.stdout.write(
                self.style.ERROR(
                    f"Contact rows with errors: "
                    f"{contact_errors}"
                )
            )

        if errors:
            self.stdout.write("")

            for error in errors:
                self.stdout.write(
                    self.style.ERROR(
                        f"  {error}"
                    )
                )

        if dry_run:
            self.stdout.write("")

            self.stdout.write(
                self.style.NOTICE(
                    "Dry run only — nothing was written "
                    "to the database."
                )
            )