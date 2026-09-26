from collections import Counter, defaultdict
from pathlib import Path

from django.contrib.contenttypes.models import ContentType
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from openpyxl import load_workbook

from outreach.models import Bank, Contact, Opportunity


IMPORT_METHOD = "Imported SRFU contact history"

ORGANISATION_COLUMNS = {
    "organisation_id",
    "bank_name",
    "state",
    "num_contacts",
    "contact_names",
    "emails",
    "contact_types",
    "any_human_reply_received",
    "all_contacts_replied",
    "overall_status",
    "first_observed_date",
    "last_observed_date",
    "needs_verification_flag",
}

CONTACT_COLUMNS = {
    "contact_id",
    "organisation_id",
    "bank_name",
    "state",
    "contact_name",
    "email",
    "contact_type",
    "name_needs_verification",
    "bank_affiliation_needs_verification",
    "first_observed_date",
    "last_observed_date",
    "sent_to_by_srfu",
    "human_reply_received",
    "contact_status",
}


def clean(value):
    if value is None:
        return ""

    if isinstance(value, float) and value.is_integer():
        return str(int(value))

    return str(value).strip()


def as_boolean(value):
    if isinstance(value, bool):
        return value

    return clean(value).lower() in {
        "true",
        "yes",
        "y",
        "1",
    }


class Command(BaseCommand):
    help = (
        "Import banks, contacts and outreach history from the cleaned "
        "SRFU Excel workbook. The command performs a dry run unless "
        "--commit is supplied."
    )

    def add_arguments(self, parser):
        parser.add_argument(
            "workbook_path",
            help="Path to the cleaned SRFU .xlsx workbook.",
        )

        mode = parser.add_mutually_exclusive_group()

        mode.add_argument(
            "--dry-run",
            action="store_true",
            help="Validate and simulate the import without saving.",
        )

        mode.add_argument(
            "--commit",
            action="store_true",
            help="Save the imported records to the database.",
        )

    def handle(self, *args, **options):
        workbook_path = Path(options["workbook_path"])
        commit_changes = options["commit"]

        if not workbook_path.exists():
            raise CommandError(
                f"Workbook not found: {workbook_path}"
            )

        if workbook_path.suffix.lower() != ".xlsx":
            raise CommandError(
                "The import file must be an .xlsx workbook."
            )

        try:
            workbook = load_workbook(
                workbook_path,
                read_only=True,
                data_only=True,
            )
        except Exception as error:
            raise CommandError(
                f"Could not open workbook: {error}"
            ) from error

        try:
            required_sheets = {
                "Organisations",
                "Contacts",
            }

            missing_sheets = (
                required_sheets - set(workbook.sheetnames)
            )

            if missing_sheets:
                raise CommandError(
                    "Workbook is missing sheet(s): "
                    + ", ".join(sorted(missing_sheets))
                )

            organisation_rows = self.read_rows(
                workbook["Organisations"],
                ORGANISATION_COLUMNS,
            )

            contact_rows = self.read_rows(
                workbook["Contacts"],
                CONTACT_COLUMNS,
            )
        finally:
            workbook.close()

        organisations = {}

        for row_number, row in organisation_rows:
            organisation_id = clean(
                row["organisation_id"]
            )

            if not organisation_id:
                raise CommandError(
                    "Organisations row "
                    f"{row_number} has no organisation_id."
                )

            if organisation_id in organisations:
                raise CommandError(
                    "Duplicate organisation_id "
                    f"'{organisation_id}' in Organisations sheet."
                )

            organisations[organisation_id] = (
                row_number,
                row,
            )

        contacts_by_organisation = defaultdict(list)

        for row_number, row in contact_rows:
            organisation_id = clean(
                row["organisation_id"]
            )

            contacts_by_organisation[
                organisation_id
            ].append(
                (row_number, row)
            )

        counts = Counter()

        self.stdout.write(
            self.style.WARNING(
                "Mode: "
                + (
                    "COMMIT — changes will be saved"
                    if commit_changes
                    else "DRY RUN — no changes will be saved"
                )
            )
        )

        with transaction.atomic():
            self.import_records(
                organisations=organisations,
                contacts_by_organisation=(
                    contacts_by_organisation
                ),
                counts=counts,
            )

            if not commit_changes:
                transaction.set_rollback(True)

        self.print_summary(
            organisations=organisations,
            contact_rows=contact_rows,
            counts=counts,
            committed=commit_changes,
        )

    @staticmethod
    def read_rows(worksheet, expected_columns):
        headers = [
            clean(cell.value).lower()
            for cell in worksheet[1]
        ]

        missing_columns = (
            expected_columns - set(headers)
        )

        if missing_columns:
            raise CommandError(
                f"{worksheet.title} is missing column(s): "
                + ", ".join(sorted(missing_columns))
            )

        rows = []

        for row_number, values in enumerate(
            worksheet.iter_rows(
                min_row=2,
                values_only=True,
            ),
            start=2,
        ):
            if not any(
                value is not None
                and clean(value) != ""
                for value in values
            ):
                continue

            row = dict(zip(headers, values))
            rows.append((row_number, row))

        return rows

    def import_records(
        self,
        *,
        organisations,
        contacts_by_organisation,
        counts,
    ):
        bank_content_type = (
            ContentType.objects.get_for_model(Bank)
        )

        imported_banks = {}

        for organisation_id, (
            row_number,
            organisation,
        ) in organisations.items():
            bank_name = clean(
                organisation["bank_name"]
            )

            state = clean(
                organisation["state"]
            )

            if not bank_name or not state:
                counts["organisations_skipped"] += 1

                self.stdout.write(
                    self.style.ERROR(
                        "Organisations row "
                        f"{row_number}: missing bank_name "
                        "or state — skipped."
                    )
                )
                continue

            source_contacts = (
                contacts_by_organisation.get(
                    organisation_id,
                    [],
                )
            )

            any_sent = any(
                as_boolean(
                    contact["sent_to_by_srfu"]
                )
                for _, contact in source_contacts
            )

            contact_status = (
                "contacted"
                if any_sent
                else "not_yet_contacted"
            )

            primary_email = next(
                (
                    clean(contact["email"])
                    for _, contact in source_contacts
                    if clean(contact["email"])
                ),
                "",
            )

            matches = list(
                Bank.objects.filter(
                    bank_name__iexact=bank_name,
                ).order_by("pk")
            )

            if len(matches) > 1:
                counts["organisations_skipped"] += 1

                self.stdout.write(
                    self.style.ERROR(
                        f"'{bank_name}' matches multiple "
                        "existing banks — skipped."
                    )
                )
                continue

            if matches:
                bank = matches[0]
                update_fields = []

                if not bank.region and state:
                    bank.region = state
                    update_fields.append("region")

                if (
                    not bank.public_email
                    and primary_email
                ):
                    bank.public_email = primary_email
                    update_fields.append(
                        "public_email"
                    )

                # Never downgrade an existing Contacted bank.
                if (
                    contact_status == "contacted"
                    and bank.contact_status
                    != "contacted"
                ):
                    bank.contact_status = "contacted"
                    update_fields.append(
                        "contact_status"
                    )

                if update_fields:
                    update_fields.append(
                        "last_updated"
                    )

                    bank.save(
                        update_fields=update_fields
                    )

                    counts[
                        "banks_updated"
                    ] += 1
                else:
                    counts[
                        "banks_unchanged"
                    ] += 1
            else:
                bank = Bank.objects.create(
                    bank_name=bank_name,
                    region=state,
                    public_email=primary_email,
                    contact_status=contact_status,
                )

                counts["banks_created"] += 1

            imported_banks[
                organisation_id
            ] = bank

            self.import_contacts(
                bank=bank,
                bank_content_type=bank_content_type,
                source_contacts=source_contacts,
                counts=counts,
            )

            if any_sent:
                self.import_opportunity(
                    bank=bank,
                    bank_content_type=(
                        bank_content_type
                    ),
                    organisation_id=(
                        organisation_id
                    ),
                    organisation=organisation,
                    source_contacts=source_contacts,
                    counts=counts,
                )

        for organisation_id, source_contacts in (
            contacts_by_organisation.items()
        ):
            if (
                organisation_id
                and organisation_id
                not in organisations
            ):
                counts[
                    "orphan_contacts"
                ] += len(source_contacts)

                self.stdout.write(
                    self.style.ERROR(
                        f"Contacts reference unknown "
                        "organisation_id "
                        f"'{organisation_id}' — skipped."
                    )
                )

    def import_contacts(
        self,
        *,
        bank,
        bank_content_type,
        source_contacts,
        counts,
    ):
        for row_number, source_contact in (
            source_contacts
        ):
            email = clean(
                source_contact["email"]
            )

            if not email:
                counts["contacts_skipped"] += 1

                self.stdout.write(
                    self.style.ERROR(
                        f"Contacts row {row_number}: "
                        "missing email — skipped."
                    )
                )
                continue

            matches = list(
                Contact.objects.filter(
                    content_type=bank_content_type,
                    object_id=bank.pk,
                    email__iexact=email,
                ).order_by("pk")
            )

            if len(matches) > 1:
                counts["contacts_skipped"] += 1

                self.stdout.write(
                    self.style.ERROR(
                        f"Contacts row {row_number}: "
                        f"multiple existing contacts use "
                        f"'{email}' for '{bank.bank_name}' "
                        "— skipped."
                    )
                )
                continue

            contact_name = clean(
                source_contact["contact_name"]
            )

            contact_type = clean(
                source_contact["contact_type"]
            )

            if matches:
                contact = matches[0]
                update_fields = []

                if (
                    not contact.contact_name
                    and contact_name
                ):
                    contact.contact_name = (
                        contact_name
                    )
                    update_fields.append(
                        "contact_name"
                    )

                if not contact.role and contact_type:
                    contact.role = contact_type
                    update_fields.append("role")

                if update_fields:
                    contact.save(
                        update_fields=update_fields
                    )
                    counts[
                        "contacts_updated"
                    ] += 1
                else:
                    counts[
                        "contacts_unchanged"
                    ] += 1
            else:
                Contact.objects.create(
                    content_type=bank_content_type,
                    object_id=bank.pk,
                    contact_name=contact_name,
                    role=contact_type,
                    email=email,
                )

                counts["contacts_created"] += 1

    def import_opportunity(
        self,
        *,
        bank,
        bank_content_type,
        organisation_id,
        organisation,
        source_contacts,
        counts,
    ):
        notes = self.build_notes(
            organisation_id=organisation_id,
            organisation=organisation,
            source_contacts=source_contacts,
        )

        matches = list(
            Opportunity.objects.filter(
                content_type=bank_content_type,
                object_id=bank.pk,
                outreach_method=IMPORT_METHOD,
            ).order_by("pk")
        )

        if len(matches) > 1:
            counts["opportunities_skipped"] += 1

            self.stdout.write(
                self.style.ERROR(
                    f"Multiple imported opportunities "
                    f"exist for '{bank.bank_name}' "
                    "— skipped."
                )
            )
            return

        if matches:
            opportunity = matches[0]
            update_fields = []

            # Preserve Interested, Not Interested and
            # Do Not Contact outcomes entered later.
            if (
                opportunity.status
                == "not_yet_contacted"
            ):
                opportunity.status = "contacted"
                update_fields.append("status")

            if opportunity.notes != notes:
                opportunity.notes = notes
                update_fields.append("notes")

            if update_fields:
                opportunity.save(
                    update_fields=update_fields
                )

                counts[
                    "opportunities_updated"
                ] += 1
            else:
                counts[
                    "opportunities_unchanged"
                ] += 1
        else:
            Opportunity.objects.create(
                content_type=bank_content_type,
                object_id=bank.pk,
                status="contacted",
                outreach_method=IMPORT_METHOD,
                notes=notes,
            )

            counts[
                "opportunities_created"
            ] += 1

    @staticmethod
    def build_notes(
        *,
        organisation_id,
        organisation,
        source_contacts,
    ):
        human_replies = sum(
            as_boolean(
                contact["human_reply_received"]
            )
            for _, contact in source_contacts
        )

        sent_contacts = sum(
            as_boolean(
                contact["sent_to_by_srfu"]
            )
            for _, contact in source_contacts
        )

        return "\n".join(
            [
                "Imported from the cleaned SRFU "
                "community-bank workbook.",
                (
                    "Source organisation ID: "
                    f"{organisation_id}"
                ),
                (
                    "Workbook response summary: "
                    f"{clean(organisation['overall_status'])}"
                ),
                (
                    "Sent contact records: "
                    f"{sent_contacts}"
                ),
                (
                    "Human replies recorded: "
                    f"{human_replies}"
                ),
                (
                    "All contacts replied: "
                    f"{as_boolean(organisation['all_contacts_replied'])}"
                ),
                (
                    "Needs verification: "
                    f"{as_boolean(organisation['needs_verification_flag'])}"
                ),
                (
                    "No Interested, Not Interested or "
                    "Do Not Contact outcome was inferred."
                ),
            ]
        )

    def print_summary(
        self,
        *,
        organisations,
        contact_rows,
        counts,
        committed,
    ):
        self.stdout.write("")
        self.stdout.write("Import summary")
        self.stdout.write(
            f"Workbook organisations: "
            f"{len(organisations)}"
        )
        self.stdout.write(
            f"Workbook contacts: {len(contact_rows)}"
        )
        self.stdout.write(
            f"Banks created: "
            f"{counts['banks_created']}"
        )
        self.stdout.write(
            f"Banks updated: "
            f"{counts['banks_updated']}"
        )
        self.stdout.write(
            f"Banks unchanged: "
            f"{counts['banks_unchanged']}"
        )
        self.stdout.write(
            f"Organisations skipped: "
            f"{counts['organisations_skipped']}"
        )
        self.stdout.write(
            f"Contacts created: "
            f"{counts['contacts_created']}"
        )
        self.stdout.write(
            f"Contacts updated: "
            f"{counts['contacts_updated']}"
        )
        self.stdout.write(
            f"Contacts unchanged: "
            f"{counts['contacts_unchanged']}"
        )
        self.stdout.write(
            f"Contacts skipped: "
            f"{counts['contacts_skipped']}"
        )
        self.stdout.write(
            f"Orphan contacts: "
            f"{counts['orphan_contacts']}"
        )
        self.stdout.write(
            f"Opportunities created: "
            f"{counts['opportunities_created']}"
        )
        self.stdout.write(
            f"Opportunities updated: "
            f"{counts['opportunities_updated']}"
        )
        self.stdout.write(
            f"Opportunities unchanged: "
            f"{counts['opportunities_unchanged']}"
        )
        self.stdout.write(
            f"Opportunities skipped: "
            f"{counts['opportunities_skipped']}"
        )

        if committed:
            self.stdout.write(
                self.style.SUCCESS(
                    "Import committed successfully."
                )
            )
        else:
            self.stdout.write(
                self.style.WARNING(
                    "Dry run complete. Nothing was "
                    "written to the database."
                )
            )