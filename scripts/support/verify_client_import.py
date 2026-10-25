"""
Read-only check that a database built from the client workbook matches the workbook.

Run from the PROJECT FOLDER with the settings/environment that point at the database to check:

    python scripts/support/verify_client_import.py "<workbook>.xlsx" [--expect-sha256 <hash>] [--expect-actor-count N]

Prints PASS/FAIL lines and counts only. It never prints organisation names, e-mail addresses or
contact names. The expected status of each organisation is computed here, independently of the
application, from the workbook's contact statuses (a confirmed outgoing e-mail makes an organisation
Contacted; delivery-failure-only, inbound-only and copied-only do not). The workbook's own
overall_status column is deliberately ignored. Exit code 0 = every check passed.
"""

import hashlib
import os
import sys
from collections import defaultdict

if len(sys.argv) < 2:
    sys.exit(__doc__)

workbook = sys.argv[1]
expect_sha = sys.argv[sys.argv.index("--expect-sha256") + 1].lower() if "--expect-sha256" in sys.argv else None
sys.path.insert(0, os.getcwd())
os.environ.setdefault("DJANGO_SETTINGS_MODULE", "srfu.settings")

import django  # noqa: E402

django.setup()

from django.conf import settings  # noqa: E402
from django.contrib.contenttypes.models import ContentType  # noqa: E402
from django.db import connection  # noqa: E402
from django.db.migrations.loader import MigrationLoader  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

from accounts.models import User  # noqa: E402
from intake.models import ContactSource, IntakeRun, OrganisationSource  # noqa: E402
from outreach.models import Bank, Branch, Club, Contact, EmailDraft, EmailTemplate, Opportunity  # noqa: E402

results = []


def check(name, ok, detail=""):
    results.append(ok)
    print(("PASS  " if ok else "FAIL  ") + name + (f"  [{detail}]" if detail else ""))


print("Database:", settings.DATABASES["default"]["NAME"])
raw = open(workbook, "rb").read()
sha = hashlib.sha256(raw).hexdigest()
print("Workbook sha256:", sha)

if expect_sha:
    check("workbook hash matches the required value", sha == expect_sha)

wb = load_workbook(workbook, read_only=True, data_only=True)
rows = list(wb["Contacts"].iter_rows(values_only=True))
header = rows[0]
ix = {name: i for i, name in enumerate(header)}
contacts = [r for r in rows[1:] if r[ix["contact_id"]] is not None]
org_rows = list(wb["Organisations"].iter_rows(values_only=True))
oix = {name: i for i, name in enumerate(org_rows[0])}
orgs_sheet = {r[oix["organisation_id"]]: r for r in org_rows[1:] if r[oix["organisation_id"]] is not None}

SENT = {"sent_and_human_reply_received", "sent_no_human_reply_found"}
by_org = defaultdict(list)

for r in contacts:
    by_org[r[ix["organisation_id"]]].append(r)

expected_contacted = {o for o, rs in by_org.items() if any(r[ix["contact_status"]] in SENT for r in rs)}
expected_blank_names = sum(1 for r in contacts if (r[ix["contact_name"]] in (None, "", "(unspecified)")))

# --- migrations / integrity --------------------------------------------------------------
loader = MigrationLoader(connection)
check("all migrations applied and none unknown", set(loader.applied_migrations) == set(loader.graph.nodes))

with connection.cursor() as cur:
    cur.execute("pragma integrity_check")
    check("sqlite integrity_check = ok", cur.fetchone()[0] == "ok")
    cur.execute("pragma foreign_key_check")
    check("no foreign-key violations", cur.fetchall() == [])

# --- counts ------------------------------------------------------------------------------
check("banks = workbook organisations", Bank.objects.count() == len(orgs_sheet), f"{Bank.objects.count()} vs {len(orgs_sheet)}")
check("contacts = workbook contact rows", Contact.objects.count() == len(contacts), f"{Contact.objects.count()} vs {len(contacts)}")
check("no Branch or Club records", Branch.objects.count() == 0 and Club.objects.count() == 0)
n_contacted = Bank.objects.filter(contact_status="contacted").count()
n_not_yet = Bank.objects.filter(contact_status="not_yet_contacted").count()
check("Contacted count equals independently derived count", n_contacted == len(expected_contacted), f"{n_contacted} vs {len(expected_contacted)}")
check("Not Yet Contacted count equals derived count", n_not_yet == len(orgs_sheet) - len(expected_contacted), f"{n_not_yet}")
check("every bank is exactly Contacted or Not Yet Contacted", n_contacted + n_not_yet == Bank.objects.count())

# --- source identifiers ------------------------------------------------------------------
bank_ct = ContentType.objects.get_for_model(Bank)
sources = list(OrganisationSource.objects.filter(source_key="client_workbook"))
src_ids = [s.source_org_id for s in sources]
check("one organisation source per bank, none duplicated", len(sources) == Bank.objects.count() and len(set(src_ids)) == len(src_ids))
check("organisation source ids = workbook organisation ids", set(src_ids) == set(orgs_sheet))
csrc = list(ContactSource.objects.filter(source_key="client_workbook").select_related("contact"))
cids = [c.source_contact_id for c in csrc]
check("one contact source per contact, none duplicated", len(csrc) == Contact.objects.count() and len(set(cids)) == len(cids))
check("contact source ids = workbook contact ids", set(cids) == {r[ix["contact_id"]] for r in contacts})

# --- each contact sits under the correct organisation -----------------------------------------
bank_by_src = {s.source_org_id: s.object_id for s in sources}
wrong = 0
contact_org = {r[ix["contact_id"]]: r[ix["organisation_id"]] for r in contacts}
for c in csrc:
    ok = c.contact.content_type_id == bank_ct.id and c.contact.object_id == bank_by_src.get(contact_org[c.source_contact_id])
    wrong += 0 if ok else 1
check("every contact is linked to the correct organisation", wrong == 0, f"{wrong} wrong")

# --- statuses per organisation (independent rule) -------------------------------------------
bad_status = 0
for src_id, bank_id in bank_by_src.items():
    want = "contacted" if src_id in expected_contacted else "not_yet_contacted"
    if Bank.objects.get(pk=bank_id).contact_status != want:
        bad_status += 1
check("each organisation's status matches the evidence rule", bad_status == 0, f"{bad_status} differ")

fail_only = [o for o, rs in by_org.items() if o not in expected_contacted and any(r[ix["contact_status"]] == "delivery_failure_recorded" for r in rs)]
check("delivery-failure-only organisations are Not Yet Contacted despite a sent flag",
      all(Bank.objects.get(pk=bank_by_src[o]).contact_status == "not_yet_contacted" for o in fail_only), f"{len(fail_only)} such organisations")

# --- names, evidence, flags --------------------------------------------------------------
check("no contact stored with the text (unspecified)", not Contact.objects.filter(contact_name__iexact="(unspecified)").exists())
check("blank contact names = workbook blank/(unspecified) names", Contact.objects.filter(contact_name="").count() == expected_blank_names,
      f"{Contact.objects.filter(contact_name='').count()} vs {expected_blank_names}")
ev_wrong = flag_wrong = 0
for c in csrc:
    r = next(x for x in contacts if x[ix["contact_id"]] == c.source_contact_id)
    ev_wrong += 0 if c.outreach_evidence == r[ix["contact_status"]] else 1
    sent = str(r[ix["sent_to_by_srfu"]]).upper() == "TRUE"
    reply = str(r[ix["human_reply_received"]]).upper() == "TRUE"
    nnv = str(r[ix["name_needs_verification"]]).upper() == "TRUE"
    flag_wrong += 0 if (c.sent_flag == sent and c.human_reply == reply and c.name_needs_verification == nnv) else 1
check("outreach evidence preserved on every contact source", ev_wrong == 0, f"{ev_wrong} differ")
check("sent / reply / verification flags preserved", flag_wrong == 0, f"{flag_wrong} differ")

# --- state and region ---------------------------------------------------------------------
state_wrong = sum(1 for src_id, bank_id in bank_by_src.items() if Bank.objects.get(pk=bank_id).state != orgs_sheet[src_id][oix["state"]])
check("State equals the workbook value for every bank", state_wrong == 0, f"{state_wrong} differ")
check("no Region values invented (all blank)", not Bank.objects.exclude(region="").exists())

# --- outcomes and contamination ---------------------------------------------------------------
statuses = set(Opportunity.objects.values_list("status", flat=True))
check("no response outcome inferred (no Interested / Not Interested / Do Not Contact)", statuses <= {"contacted", "not_yet_contacted"}, str(sorted(statuses)))
check("no email drafts", EmailDraft.objects.count() == 0)
check("no email templates", EmailTemplate.objects.count() == 0)
check("no scraper jobs", IntakeRun.objects.filter(kind="scraper").count() == 0)
check("no archived or manually created banks", not Bank.objects.exclude(record_source="Client Import").exists() and not Bank.objects.filter(is_archived=True).exists())

with connection.cursor() as cur:
    cur.execute("select count(*) from django_session")
    sessions = cur.fetchone()[0]

check("no login sessions", sessions == 0, str(sessions))
n_users = User.objects.count()
expected_users = int(sys.argv[sys.argv.index("--expect-actor-count") + 1]) if "--expect-actor-count" in sys.argv else None
print("INFO  users in database:", n_users, "| import runs:", IntakeRun.objects.filter(kind="import").count())

if expected_users is not None:
    check(f"user count = {expected_users}", n_users == expected_users)

print("RESULT:", "ALL CHECKS PASSED" if all(results) else "SOME CHECKS FAILED", f"({sum(results)}/{len(results)})")
sys.exit(0 if all(results) else 1)
