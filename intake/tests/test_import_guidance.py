"""Import-format guidance, the downloadable template and invalid-upload safety (synthetic data only)."""

import io

from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
from openpyxl import Workbook, load_workbook

from intake.import_template import (
    CLIENT_EXAMPLE_FILENAME,
    CLIENT_TEMPLATE_FILENAME,
    EXAMPLE_FILENAME,
    OPTIONAL_ORGANISATION_HEADERS,
    TEMPLATE_FILENAME,
    build_client_example,
    build_client_template,
    build_example,
    build_template,
)
from intake.importer import confirm_import, create_preview
from intake.models import IntakeRun
from intake.workbook import CONTACT_HEADERS, ORGANISATION_HEADERS, SIMPLE_HEADERS, parse_workbook
from outreach.models import Bank, Branch, Club, Contact

from .helpers import IntakeTestCase, con_row, make_workbook, org_row

URLS = (
    "intake_import_template", "intake_import_example",
    "intake_import_client_template", "intake_import_client_example",
)


def nothing_imported():
    return (Bank.objects.count(), Branch.objects.count(), Club.objects.count(), Contact.objects.count()) == (0, 0, 0, 0)


class TemplateDownloadTests(IntakeTestCase):
    def test_only_admins_can_download(self):
        for name in URLS:
            url = reverse(name)
            self.assertEqual(self.client.get(url).status_code, 302, "anonymous users are sent to log in")
            self.client.force_login(self.staff)
            self.assertEqual(self.client.get(url).status_code, 403, "Staff must not download import files")
            self.client.logout()
            self.client.force_login(self.admin)
            response = self.client.get(url)
            self.assertEqual(response.status_code, 200)
            self.assertIn("spreadsheetml.sheet", response["Content-Type"])
            self.assertIn("attachment", response["Content-Disposition"])
            self.client.logout()

    def test_file_names(self):
        self.client.force_login(self.admin)

        self.assertIn(TEMPLATE_FILENAME, self.client.get(reverse("intake_import_template"))["Content-Disposition"])
        self.assertIn(EXAMPLE_FILENAME, self.client.get(reverse("intake_import_example"))["Content-Disposition"])
        self.assertIn(CLIENT_TEMPLATE_FILENAME, self.client.get(reverse("intake_import_client_template"))["Content-Disposition"])
        self.assertIn(CLIENT_EXAMPLE_FILENAME, self.client.get(reverse("intake_import_client_example"))["Content-Disposition"])
        self.assertIn("EXAMPLE", EXAMPLE_FILENAME)
        self.assertIn("synthetic", EXAMPLE_FILENAME)

    def test_template_is_the_single_sheet_minimum_columns_with_no_data(self):
        wb = load_workbook(io.BytesIO(build_template()))

        self.assertEqual(wb.sheetnames, ["Organisations", "Instructions"])
        self.assertEqual([c.value for c in wb["Organisations"][1]], list(SIMPLE_HEADERS))
        self.assertEqual(wb["Organisations"].max_row, 1, "headers only")
        self.assertEqual(
            list(SIMPLE_HEADERS),
            ["organisation_id", "bank_name", "state", "contact_names", "emails", "overall_status"],
        )
        everything = " ".join(str(c.value) for sheet in wb for row in sheet.iter_rows() for c in row if c.value)
        self.assertIn("never pairs names and emails", everything)
        self.assertNotIn("@", everything)

    def test_client_template_has_the_supported_structure_and_no_data(self):
        wb = load_workbook(io.BytesIO(build_client_template()))

        self.assertEqual(wb.sheetnames[:2], ["Organisations", "Contacts"])
        organisations, contacts = wb["Organisations"], wb["Contacts"]
        self.assertEqual([c.value for c in organisations[1]], list(ORGANISATION_HEADERS + OPTIONAL_ORGANISATION_HEADERS))
        self.assertEqual([c.value for c in contacts[1]], list(CONTACT_HEADERS))
        self.assertEqual((organisations.max_row, contacts.max_row), (1, 1), "headers only")
        everything = " ".join(str(c.value) for sheet in wb for row in sheet.iter_rows() for c in row if c.value)
        for forbidden in ("@", "password", "secret", "api_key"):
            self.assertNotIn(forbidden, everything.lower())

    def test_every_header_in_the_template_is_one_the_parser_requires_or_accepts(self):
        # Fill one organisation into the client template headers: it must parse.
        wb = load_workbook(io.BytesIO(build_client_template()))
        wb["Organisations"].append(["T_1", "Template Test Bank", "VIC"] + [""] * (len(ORGANISATION_HEADERS) - 3) + ["Bank"] + [""] * (len(OPTIONAL_ORGANISATION_HEADERS) - 1))
        buffer = io.BytesIO()
        wb.save(buffer)

        parsed = parse_workbook(buffer.getvalue(), "t.xlsx")

        self.assertTrue(parsed.ok, parsed.structure_errors)
        self.assertEqual(parsed.groups[0]["org_type"], "bank")

    def test_blank_templates_cannot_be_previewed_and_say_why(self):
        for build in (build_template, build_client_template):
            with self.subTest(build=build.__name__):
                run, parsed = create_preview(self.admin, build(), "blank.xlsx")

                self.assertFalse(parsed.ok)
                self.assertIn("headers-only", parsed.structure_errors[0])
                self.assertIn("add one row per organisation", parsed.structure_errors[0])
                self.assertEqual(run.status, IntakeRun.STATUS_FAILED)
                self.assertTrue(nothing_imported())

    def test_filled_single_sheet_template_parses(self):
        wb = load_workbook(io.BytesIO(build_template()))
        wb["Organisations"].append(["T_1", "Template Test Bank", "VIC", "", "", ""])
        buffer = io.BytesIO()
        wb.save(buffer)

        parsed = parse_workbook(buffer.getvalue(), "t.xlsx")

        self.assertTrue(parsed.ok, parsed.structure_errors)
        self.assertEqual(parsed.format, "simple")
        self.assertEqual(parsed.columns[0]["unsupported"], [])

    def test_single_sheet_example_previews_and_imports_then_repeats_safely(self):
        run, parsed = create_preview(self.admin, build_example(), "example.xlsx")
        self.assertTrue(parsed.ok, parsed.structure_errors)
        self.assertEqual(parsed.row_errors, [])
        self.assertEqual(parsed.columns[0]["unsupported"], ["internal_notes"])
        self.assertTrue(nothing_imported(), "the preview must not create records")

        confirm_import(run.pk, self.admin)

        self.assertEqual((Bank.objects.count(), Club.objects.count(), Contact.objects.count()), (3, 1, 5))
        self.assertEqual({b.contact_status for b in Bank.objects.all()}, {"not_yet_contacted"})
        self.assertEqual(Contact.objects.get(email="pat@example.test").contact_name, "Pat Example")
        self.assertFalse(Contact.objects.exclude(email="pat@example.test").exclude(contact_name="").exists())

        again, _ = create_preview(self.admin, build_example(), "example.xlsx")
        confirm_import(again.pk, self.admin)

        self.assertEqual((Bank.objects.count(), Club.objects.count(), Contact.objects.count()), (3, 1, 5))

    def test_client_example_previews_and_imports_then_repeats_safely(self):
        run, parsed = create_preview(self.admin, build_client_example(), "example.xlsx")
        self.assertTrue(parsed.ok, parsed.structure_errors)
        self.assertEqual(parsed.format, "client")
        self.assertEqual(parsed.row_errors, [])
        self.assertTrue(nothing_imported(), "the preview must not create records")

        confirm_import(run.pk, self.admin)
        run.refresh_from_db()

        self.assertEqual(run.status, IntakeRun.STATUS_SUCCESSFUL)
        self.assertEqual((Bank.objects.count(), Club.objects.count(), Contact.objects.count()), (2, 1, 3))
        alpha = Bank.objects.get(bank_name__contains="Alpha")
        self.assertEqual(alpha.contact_status, "contacted")
        self.assertEqual(Bank.objects.get(bank_name__contains="Bravo").contact_status, "not_yet_contacted")
        self.assertEqual(Club.objects.get().club_type, "Netball")
        self.assertTrue(all("EXAMPLE ONLY" in b.bank_name for b in Bank.objects.all()))
        self.assertEqual(Contact.objects.get(email="info@example.test").contact_name, "")

        again, _ = create_preview(self.admin, build_client_example(), "example.xlsx")
        confirm_import(again.pk, self.admin)

        self.assertEqual((Bank.objects.count(), Club.objects.count(), Contact.objects.count()), (2, 1, 3))


class InvalidUploadTests(IntakeTestCase):
    def post(self, content, name):
        self.client.force_login(self.admin)

        return self.client.post(
            reverse("intake_import"),
            {"workbook": SimpleUploadedFile(name, content)},
        )

    def test_invalid_uploads_name_the_problem_and_create_nothing(self):
        only_contacts = make_workbook([], [con_row()], sheets=("Contacts",))
        thin = make_workbook([org_row()], [], org_header=["organisation_id", "bank_name", "state"], sheets=("Organisations",))
        lower = Workbook()
        lower.remove(lower.active)
        lower.create_sheet("organisations")
        lower.create_sheet("contacts")
        buffer = io.BytesIO()
        lower.save(buffer)
        missing_cols = make_workbook([org_row()], [con_row()], con_header=["contact_id", "organisation_id", "email"])
        cases = [
            (b"name,state\nA,VIC\n", "list.csv", "Only Microsoft Excel"),
            (b"name,state\nA,VIC\n", "renamed.xlsx", "cannot be read"),
            (only_contacts, "no-orgs.xlsx", "no &#x27;Organisations&#x27; sheet"),
            (thin, "thin.xlsx", "missing the column(s): contact_names, emails, overall_status"),
            (buffer.getvalue(), "case.xlsx", "Sheets found: &#x27;organisations&#x27;, &#x27;contacts&#x27;"),
            (missing_cols, "cols.xlsx", "missing the column(s): bank_affiliation_needs_verification"),
        ]

        for content, name, expected in cases:
            with self.subTest(name=name):
                response = self.post(content, name)

                self.assertEqual(response.status_code, 200, "no redirect to a preview")
                self.assertContains(response, expected)
                self.assertTrue(nothing_imported())

        self.assertFalse(IntakeRun.objects.filter(status__in=("previewed", "running", "successful", "partial")).exists())

    def test_the_csv_message_explains_renaming(self):
        response = self.post(b"a,b\n", "renamed.xlsx")

        self.assertContains(response, "only been renamed to .xlsx")
        self.assertContains(response, "Excel Workbook (.xlsx)")

    def test_a_valid_upload_still_goes_to_preview_then_confirm(self):
        response = self.post(build_client_example(), "example.xlsx")

        self.assertEqual(response.status_code, 302)
        run = IntakeRun.objects.get()
        self.assertEqual(run.status, IntakeRun.STATUS_PREVIEWED)
        self.assertTrue(nothing_imported(), "nothing is saved before confirmation")


class GuidancePageTests(IntakeTestCase):
    def test_page_is_short_and_offers_the_template_next_to_the_upload(self):
        self.client.force_login(self.admin)

        page = self.client.get(reverse("intake_import"))

        self.assertContains(page, "cannot simply be renamed to .xlsx")
        self.assertContains(page, reverse("intake_import_template"))
        self.assertContains(page, reverse("intake_import_example"))
        self.assertContains(page, reverse("intake_import_client_template"))
        self.assertContains(page, "headers only")
        self.assertContains(page, "two-sheet client workbook")
        self.assertNotContains(page, "Required import format")
        self.assertNotContains(page, "Show the detailed instructions")
        self.assertLess(len(page.content.decode()), 12000)

    def test_page_names_the_minimum_columns_and_the_status_rule(self):
        self.client.force_login(self.admin)
        page = self.client.get(reverse("intake_import")).content.decode()

        for header in SIMPLE_HEADERS:
            self.assertIn(f"<code>{header}</code>", page, header)

        self.assertIn("never marks an organisation Contacted on its own", page)
        self.assertNotIn("Partially Contacted</td>", page)

    def test_staff_cannot_see_the_page_or_its_downloads(self):
        self.client.force_login(self.staff)

        self.assertEqual(self.client.get(reverse("intake_import")).status_code, 403)
