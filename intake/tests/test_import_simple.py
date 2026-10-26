"""Single-sheet import format (synthetic data only); the two-sheet format is covered in test_import.py."""

import io

from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
from openpyxl import Workbook

from intake.importer import ImportBusy, confirm_import, create_preview
from intake.models import ContactSource, DuplicateReview, IntakeResult, IntakeRun, OrganisationSource
from intake.reviews import decide_review
from intake.workbook import SIMPLE_HEADERS, contacts_from_cells, parse_workbook
from outreach.models import Bank, Club, Contact, Opportunity

from .helpers import IntakeTestCase, con_row, make_workbook, org_row


def simple_workbook(rows, header=SIMPLE_HEADERS, sheet="Organisations", extra_sheets=()):
    wb = Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet(sheet)
    ws.append(list(header))

    for row in rows:
        ws.append(list(row))

    for name in extra_sheets:
        wb.create_sheet(name).append(["ignored"])

    buffer = io.BytesIO()
    wb.save(buffer)

    return buffer.getvalue()


def row(oid="S_1", name="Alpha Community Bank", state="VIC", names="", emails="", status=""):
    return [oid, name, state, names, emails, status]


def do_import(user, data, name="simple.xlsx"):
    run, parsed = create_preview(user, data, name)

    if not parsed.ok:
        return run, parsed

    return confirm_import(run.pk, user), parsed


class MinimumColumnTests(IntakeTestCase):
    def test_minimum_columns_import_with_one_name_and_one_email_paired(self):
        run, parsed = do_import(self.admin, simple_workbook([
            row("S_1", "Alpha Community Bank", "VIC", "Pat Example", "pat@example.test", "Not Yet Contacted"),
        ]))

        self.assertTrue(parsed.ok)
        self.assertEqual(parsed.format, "simple")
        self.assertEqual(run.status, IntakeRun.STATUS_SUCCESSFUL)
        bank = Bank.objects.get()
        self.assertEqual((bank.bank_name, bank.state, bank.contact_status), ("Alpha Community Bank", "VIC", "not_yet_contacted"))
        self.assertEqual(Contact.objects.get().contact_name, "Pat Example")
        self.assertEqual(Contact.objects.get().email, "pat@example.test")

    def test_source_id_and_the_files_own_status_are_kept_for_reference(self):
        do_import(self.admin, simple_workbook([row("S_9", "Alpha Community Bank", "VIC", "", "a@example.test", "Not Yet Contacted")]),
                  "simple.xlsx")

        source = OrganisationSource.objects.get()
        self.assertEqual(source.source_org_id, "S_9")
        self.assertEqual(source.source_filename, "simple.xlsx")
        self.assertEqual(source.extra["overall_status_in_file"], "Not Yet Contacted")
        self.assertEqual(source.extra["import_format"], "simple")
        self.assertEqual(source.extra["emails_in_file"], "a@example.test")

    def test_contact_names_emails_and_status_may_be_blank(self):
        run, parsed = do_import(self.admin, simple_workbook([row("S_1", "Alpha Community Bank", "VIC")]))

        self.assertEqual(parsed.row_errors, [])
        self.assertEqual(run.status, IntakeRun.STATUS_SUCCESSFUL)
        self.assertEqual(Bank.objects.get().contact_status, "not_yet_contacted")
        self.assertEqual(Contact.objects.count(), 0)

    def test_id_name_and_state_are_required_and_other_rows_continue(self):
        run, parsed = do_import(self.admin, simple_workbook([
            row("", "No Id Bank", "VIC"),
            row("S_2", "", "VIC"),
            row("S_3", "No State Bank", ""),
            row("S_4", "Odd State Bank", "Narnia"),
            row("S_5", "Good Community Bank", "SA"),
        ]))

        self.assertEqual([e["field"] for e in parsed.row_errors], ["organisation_id", "bank_name", "state", "state"])
        self.assertEqual(Bank.objects.get().bank_name, "Good Community Bank")
        self.assertEqual(run.totals["created"], 1)
        self.assertEqual(IntakeResult.objects.filter(run=run, outcome="failed").count(), 4)

    def test_repeated_organisation_id_is_reported(self):
        _run, parsed = create_preview(self.admin, simple_workbook([row("S_1"), row("S_1", "Other Community Bank")]), "s.xlsx")

        self.assertEqual(len(parsed.groups), 1)
        self.assertIn("more than once", parsed.row_errors[0]["message"])

    def test_missing_minimum_header_is_named_and_nothing_is_saved(self):
        run, parsed = create_preview(
            self.admin, simple_workbook([["S_1", "A Bank", "VIC"]], header=("organisation_id", "bank_name", "state")), "s.xlsx",
        )

        self.assertFalse(parsed.ok)
        self.assertIn("contact_names, emails, overall_status", parsed.structure_errors[0])
        self.assertEqual(run.status, IntakeRun.STATUS_FAILED)
        self.assertEqual(Bank.objects.count(), 0)

    def test_a_repeated_header_is_refused_rather_than_guessed(self):
        _run, parsed = create_preview(
            self.admin, simple_workbook([row()], header=SIMPLE_HEADERS + ("emails",)), "s.xlsx",
        )

        self.assertFalse(parsed.ok)
        self.assertIn("emails more than once", parsed.structure_errors[0])


class SheetSelectionTests(IntakeTestCase):
    def test_a_single_sheet_with_any_name_is_accepted(self):
        parsed = parse_workbook(simple_workbook([row()], sheet="Banks"), "s.xlsx")

        self.assertTrue(parsed.ok)
        self.assertEqual(parsed.groups[0]["row_ref"], "Banks row 2")

    def test_organisations_sheet_wins_and_other_sheets_including_instructions_are_ignored(self):
        parsed = parse_workbook(simple_workbook([row()], extra_sheets=("Instructions", "Notes")), "s.xlsx")

        self.assertTrue(parsed.ok)
        self.assertEqual(parsed.format, "simple")

    def test_several_unnamed_sheets_are_refused(self):
        wb = Workbook()
        wb.active.title = "One"
        wb.create_sheet("Two")
        buffer = io.BytesIO()
        wb.save(buffer)

        parsed = parse_workbook(buffer.getvalue(), "s.xlsx")

        self.assertIn("none is named 'Organisations'", parsed.structure_errors[0])

    def test_both_sheets_still_mean_the_client_workbook(self):
        parsed = parse_workbook(make_workbook([org_row()], [con_row()]), "c.xlsx")

        self.assertTrue(parsed.ok)
        self.assertEqual(parsed.format, "client")
        self.assertEqual(parsed.groups[0]["evidence_status"], "contacted")


class ColumnReportTests(IntakeTestCase):
    def test_imported_informational_and_unsupported_columns_are_identified(self):
        header = SIMPLE_HEADERS + ("suburb", "postcode", "organisation_type", "num_contacts", "internal_notes", "Colour")
        parsed = parse_workbook(simple_workbook([row() + ["Exampleton", "3999", "Club", 1, "x", "y"]], header=header), "s.xlsx")

        columns = parsed.columns[0]
        self.assertEqual(columns["imported"], list(SIMPLE_HEADERS) + ["suburb", "postcode", "organisation_type"])
        self.assertEqual(columns["informational"], ["num_contacts"])
        self.assertEqual(columns["unsupported"], ["internal_notes", "colour"])

    def test_additional_supported_columns_are_imported(self):
        header = SIMPLE_HEADERS + ("organisation_type", "suburb", "postcode", "club_type")
        do_import(self.admin, simple_workbook([row("S_1", "Alpha Netball Club", "VIC") + ["Club", "Exampleton", "3999", "Netball"]], header=header))

        club = Club.objects.get()
        self.assertEqual((club.suburb, club.postcode, club.club_type), ("Exampleton", "3999", "Netball"))

    def test_preview_page_lists_columns_and_flags_unsupported_ones(self):
        header = SIMPLE_HEADERS + ("internal_notes",)
        self.client.force_login(self.admin)

        response = self.client.post(
            reverse("intake_import"),
            {"workbook": SimpleUploadedFile("s.xlsx", simple_workbook([row() + ["x"]], header=header))},
            follow=True,
        )

        self.assertContains(response, "Single-sheet workbook")
        self.assertContains(response, "will be imported")
        self.assertContains(response, "Unsupported columns on Organisations")
        self.assertContains(response, "<code>internal_notes</code>")
        self.assertEqual(Bank.objects.count(), 0, "nothing is saved until confirmation")

    def test_client_workbook_unsupported_columns_are_flagged_too(self):
        data = make_workbook(
            [org_row(extra_thing="x")], [con_row()],
            org_header=list(__import__("intake.tests.helpers", fromlist=["ORG_HEADER"]).ORG_HEADER) + ["extra_thing"],
        )
        parsed = parse_workbook(data, "c.xlsx")

        self.assertEqual(parsed.columns[0]["unsupported"], ["extra_thing"])
        self.assertEqual(parsed.columns[1]["unsupported"], [])
        self.assertIn("contact_names", parsed.columns[0]["informational"])


class MultipleContactTests(IntakeTestCase):
    def pair(self, names, emails):
        return contacts_from_cells("Row 2", names, emails)

    def test_one_name_and_one_email_are_paired(self):
        contacts, notes = self.pair("Pat Example", "pat@example.test")

        self.assertEqual([(c["name"], c["email"]) for c in contacts], [("Pat Example", "pat@example.test")])
        self.assertEqual(notes, [])

    def test_a_name_with_a_comma_is_one_name(self):
        contacts, notes = self.pair("Example, Pat", "pat@example.test")

        self.assertEqual(contacts[0]["name"], "Example, Pat")
        self.assertEqual(notes, [])

    def test_several_emails_without_names_need_no_review(self):
        contacts, notes = self.pair("", "a@example.test; b@example.test\nc@example.test, A@example.test")

        self.assertEqual([c["email"] for c in contacts], ["a@example.test", "b@example.test", "c@example.test"])
        self.assertEqual({c["name"] for c in contacts}, {""})
        self.assertEqual(notes, [])

    def test_matching_counts_are_not_paired_by_position(self):
        contacts, notes = self.pair("Sam One; Sam Two", "one@example.test; two@example.test")

        self.assertEqual({c["name"] for c in contacts}, {""})
        self.assertEqual(len(contacts), 2)
        self.assertEqual([n["outcome"] for n in notes], ["review"])
        self.assertIn("Sam One; Sam Two", notes[0]["message"])

    def test_one_name_with_several_emails_and_several_names_with_one_email_are_ambiguous(self):
        for names, emails in (("Sam One", "one@example.test; two@example.test"), ("Sam One; Sam Two", "one@example.test")):
            with self.subTest(names=names):
                contacts, notes = self.pair(names, emails)

                self.assertEqual({c["name"] for c in contacts}, {""})
                self.assertEqual([n["outcome"] for n in notes], ["review"])

    def test_names_without_any_email_create_nothing(self):
        contacts, notes = self.pair("Sam One", "")

        self.assertEqual(contacts, [])
        self.assertEqual([n["outcome"] for n in notes], ["failed"])

    def test_an_invalid_email_skips_only_that_contact(self):
        contacts, notes = self.pair("", "good@example.test; not-an-email")

        self.assertEqual([c["email"] for c in contacts], ["good@example.test"])
        self.assertEqual([n["field"] for n in notes], ["emails"])

    def test_placeholder_names_are_blank(self):
        contacts, notes = self.pair("(unspecified)", "a@example.test")

        self.assertEqual(contacts[0]["name"], "")
        self.assertEqual(notes, [])

    def test_ambiguous_row_is_flagged_in_the_preview_and_in_the_results(self):
        run, parsed = create_preview(self.admin, simple_workbook([
            row("S_1", "Alpha Community Bank", "VIC", "Sam One; Sam Two", "one@example.test; two@example.test"),
        ]), "s.xlsx")

        self.assertEqual(run.totals["preview"]["contact_review"], 1)
        self.assertEqual(Contact.objects.count(), 0, "the preview writes nothing")

        confirm_import(run.pk, self.admin)
        run.refresh_from_db()

        self.assertEqual(sorted(Contact.objects.values_list("email", "contact_name")),
                         [("one@example.test", ""), ("two@example.test", "")])
        result = IntakeResult.objects.get(run=run, kind="contact")
        self.assertEqual((result.outcome, result.field), ("review", "contact_names"))
        self.assertEqual(run.status, IntakeRun.STATUS_PARTIAL)
        self.assertEqual(run.totals["review"], 1)
        self.assertEqual(OrganisationSource.objects.get().extra["contact_names_in_file"], "Sam One; Sam Two")


class StatusEvidenceTests(IntakeTestCase):
    def test_a_contacted_claim_alone_is_held_for_review_and_not_applied(self):
        for status in ("Contacted", "contacted", "Partially Contacted"):
            with self.subTest(status=status):
                Bank.objects.all().delete()
                DuplicateReview.objects.all().delete()
                run, _ = do_import(self.admin, simple_workbook([row("S_1", "Alpha Community Bank", "VIC", "", "a@example.test", status)]), f"{status}.xlsx")

                self.assertEqual(Bank.objects.count(), 0)
                self.assertEqual(Opportunity.objects.count(), 0)
                review = DuplicateReview.objects.get()
                self.assertEqual(review.kind, DuplicateReview.KIND_EVIDENCE)
                self.assertIn("no outreach evidence", review.reason)
                self.assertEqual(run.status, IntakeRun.STATUS_PARTIAL)

    def test_an_unrecognised_status_is_held_not_guessed(self):
        do_import(self.admin, simple_workbook([row("S_1", "Alpha Community Bank", "VIC", "", "", "Maybe emailed once")]))

        self.assertEqual(Bank.objects.count(), 0)
        self.assertIn("not recognised", DuplicateReview.objects.get().reason)

    def test_blank_and_not_yet_contacted_are_imported_as_not_yet_contacted(self):
        do_import(self.admin, simple_workbook([
            row("S_1", "Alpha Community Bank", "VIC", status=""),
            row("S_2", "Bravo Community Bank", "NSW", status="not_yet_contacted"),
        ]))

        self.assertEqual({b.contact_status for b in Bank.objects.all()}, {"not_yet_contacted"})

    def test_an_admin_may_confirm_a_held_claim(self):
        do_import(self.admin, simple_workbook([row("S_1", "Alpha Community Bank", "VIC", "", "", "Contacted")]))

        decide_review(DuplicateReview.objects.get().pk, self.admin, "accept_contacted")

        self.assertEqual(Bank.objects.get().contact_status, "contacted")
        self.assertIn("Admin", Opportunity.objects.get().notes)

    def test_reply_flags_alone_never_set_a_status(self):
        header = SIMPLE_HEADERS + ("any_human_reply_received",)
        do_import(self.admin, simple_workbook([row("S_1", "Alpha Community Bank", "VIC") + ["TRUE"]], header=header))

        self.assertEqual(Bank.objects.get().contact_status, "not_yet_contacted")
        self.assertEqual(Opportunity.objects.count(), 0)
        self.assertIs(OrganisationSource.objects.get().extra["any_human_reply"], True)


class DuplicateAndRepeatTests(IntakeTestCase):
    def test_an_existing_exact_organisation_is_matched_and_blank_fields_filled(self):
        existing = self.bank("Alpha Community Bank", "VIC", region="")

        run, _ = do_import(self.admin, simple_workbook([row("S_1", "Alpha Community Bank", "VIC", "", "a@example.test")]))

        self.assertEqual(Bank.objects.count(), 1)
        self.assertEqual(run.totals["exact_duplicates"], 1)
        self.assertEqual(Contact.objects.get().object_id, existing.pk)

    def test_a_similar_name_is_held_for_duplicate_review(self):
        self.bank("Alpha Community Bank", "VIC")

        do_import(self.admin, simple_workbook([row("S_1", "Alpha Community Bank Ltd Branch", "VIC")]))

        self.assertEqual(Bank.objects.count(), 1)
        self.assertEqual(DuplicateReview.objects.get().kind, DuplicateReview.KIND_DUPLICATE)

    def test_repeat_import_creates_nothing_new_and_the_page_warns(self):
        data = simple_workbook([row("S_1", "Alpha Community Bank", "VIC", "Pat Example", "pat@example.test")])
        do_import(self.admin, data)
        before = (Bank.objects.count(), Contact.objects.count())

        run, _ = create_preview(self.admin, data, "simple.xlsx")
        self.client.force_login(self.admin)
        self.assertContains(self.client.get(reverse("intake_run", args=[run.pk])), "already imported")
        confirm_import(run.pk, self.admin)

        self.assertEqual((Bank.objects.count(), Contact.objects.count()), before)
        run.refresh_from_db()
        self.assertEqual(run.totals["created"], 0)
        self.assertEqual(run.totals["contacts_created"], 0)

    def test_a_preview_confirms_once_only(self):
        run, _ = create_preview(self.admin, simple_workbook([row()]), "s.xlsx")
        confirm_import(run.pk, self.admin)

        with self.assertRaises(ImportBusy):
            confirm_import(run.pk, self.admin)

        self.assertEqual(Bank.objects.count(), 1)

    def test_a_different_organisation_sharing_an_email_is_held_and_not_queued_twice(self):
        do_import(self.admin, simple_workbook([row("S_1", "Alpha Community Bank", "VIC", "", "shared@example.test")]))

        do_import(self.admin, simple_workbook([row("S_2", "Zulu Community Bank", "NSW", "", "shared@example.test")]), "other.xlsx")
        do_import(self.admin, simple_workbook([row("S_2", "Zulu Community Bank", "NSW", "", "shared@example.test")]), "other-again.xlsx")

        self.assertEqual((Bank.objects.count(), Contact.objects.count()), (1, 1))
        review = DuplicateReview.objects.get()
        self.assertEqual((review.kind, review.status), (DuplicateReview.KIND_DUPLICATE, "pending"), "a repeat must not queue it twice")


class PreservationTests(IntakeTestCase):
    """A later single-sheet file must not erase what an earlier two-sheet import kept."""

    def test_two_sheet_evidence_and_history_survive_a_single_sheet_import(self):
        do_import(self.admin, make_workbook([org_row("O_1", "Alpha Community Bank", "VIC")], [con_row("C_1", "O_1", "alpha@example.test", name="Pat Example")]), "client.xlsx")
        bank = Bank.objects.get()
        self.assertEqual(bank.contact_status, "contacted")

        do_import(self.admin, simple_workbook([row("O_1", "Alpha Community Bank", "VIC", "Pat Example", "alpha@example.test", "Not Yet Contacted")]), "simple.xlsx")

        bank.refresh_from_db()
        self.assertEqual(bank.contact_status, "contacted", "never downgraded")
        self.assertEqual(Bank.objects.count(), 1)
        self.assertEqual(Contact.objects.count(), 1)
        cs = ContactSource.objects.get()
        self.assertEqual(cs.source_contact_id, "C_1")
        self.assertEqual(cs.outreach_evidence, "sent_no_human_reply_found")
        self.assertIs(cs.sent_flag, True)
        self.assertIs(cs.name_needs_verification, True)
        self.assertIsNotNone(cs.first_observed_at)
        source = OrganisationSource.objects.get()
        self.assertEqual(source.source_org_id, "O_1")
        self.assertEqual(source.extra["first_observed"][:10], "2025-01-10")
        self.assertIs(source.extra["needs_verification_flag"], True)
        self.assertTrue(source.requires_verification)
        self.assertEqual(Opportunity.objects.count(), 1)


class AccessTests(IntakeTestCase):
    def test_only_admins_can_upload_preview_or_confirm(self):
        data = simple_workbook([row()])
        self.client.force_login(self.staff)

        self.assertEqual(self.client.post(reverse("intake_import"), {"workbook": SimpleUploadedFile("s.xlsx", data)}).status_code, 403)
        self.assertEqual(Bank.objects.count(), 0)
        run, _ = create_preview(self.admin, data, "s.xlsx")
        self.assertEqual(self.client.post(reverse("intake_confirm", args=[run.pk])).status_code, 403)
        run.refresh_from_db()
        self.assertEqual(run.status, IntakeRun.STATUS_PREVIEWED)

        with self.assertRaises(Exception):
            create_preview(self.staff, data, "s.xlsx")
