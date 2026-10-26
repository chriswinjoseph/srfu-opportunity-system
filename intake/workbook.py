"""
Read and validate the agreed client workbook (.xlsx) before anything is
written. Nothing here touches the database.
"""

import hashlib
import io
import re
from collections import defaultdict
from datetime import date, datetime
from zoneinfo import ZoneInfo

from django.core.validators import validate_email
from django.core.exceptions import ValidationError
from openpyxl import load_workbook

from .evidence import map_claimed_status, map_organisation_status
from .normalise import clean_contact_name, collapse, norm_state, parse_bool

MELBOURNE = ZoneInfo("Australia/Melbourne")
MAX_BYTES = 8 * 1024 * 1024
MAX_ROWS = 20000

SOURCE_KEY = "client_workbook"
SOURCE_LABEL = "Client workbook (Shane)"

FORMAT_SIMPLE = "simple"
FORMAT_CLIENT = "client"
FORMAT_LABELS = {
    FORMAT_SIMPLE: "Single-sheet workbook",
    FORMAT_CLIENT: "Two-sheet client workbook",
}

SHEET_ORGANISATIONS = "Organisations"
SHEET_CONTACTS = "Contacts"

# Ordered so the downloadable template and the on-page guidance list the
# headers in the same order the sheets use. The sets below are what is checked.
ORGANISATION_HEADERS = (
    "organisation_id",
    "bank_name",
    "state",
    "num_contacts",
    "contact_names",
    "emails",
    "contact_types",
    "any_human_reply_received",
    "all_contacts_replied",
    "overall_status",
    "first_observed_date",
    "last_observed_date",
    "needs_verification_flag",
)
CONTACT_HEADERS = (
    "contact_id",
    "organisation_id",
    "bank_name",
    "state",
    "contact_name",
    "email",
    "contact_type",
    "name_needs_verification",
    "bank_affiliation_needs_verification",
    "first_observed_date",
    "last_observed_date",
    "sent_to_by_srfu",
    "human_reply_received",
    "contact_status",
)
# The single-sheet format needs only these headers. contact_names, emails and
# overall_status may be blank; organisation_id, bank_name and state may not.
SIMPLE_HEADERS = (
    "organisation_id",
    "bank_name",
    "state",
    "contact_names",
    "emails",
    "overall_status",
)
ORGANISATION_COLUMNS = set(ORGANISATION_HEADERS)
CONTACT_COLUMNS = set(CONTACT_HEADERS)
TYPE_COLUMN = "organisation_type"

SUPPORTED_TYPES = {"bank": "bank", "branch": "branch", "club": "club"}

# Optional organisation columns for future (non-bank) imports. They are read
# only when the column exists; blank cells stay empty.
OPTIONAL_COLUMNS = {
    "parent_bank_name": "parent_bank_name",
    "parent_bank_state": "parent_bank_state",
    "region": "region",
    "suburb": "suburb",
    "postcode": "postcode",
    "address": "address",
    "website_url": "website",
    "public_phone": "phone",
    "public_email": "email",
    "club_type": "club_type",
}

# Columns that are read and stored, per sheet. A header that is in
# ORGANISATION_HEADERS / CONTACT_HEADERS but not stored is recognised and
# informational only; every other header is unsupported and ignored (and is
# shown as such in the preview).
_ORG_STORED = {
    "organisation_id", "bank_name", "state", "overall_status", "any_human_reply_received",
    "all_contacts_replied", "first_observed_date", "last_observed_date",
    "needs_verification_flag", TYPE_COLUMN,
} | set(OPTIONAL_COLUMNS)
SIMPLE_STORED = _ORG_STORED | {"contact_names", "emails"}
SIMPLE_INFORMATIONAL = {"num_contacts", "contact_types"}
CLIENT_ORG_STORED = _ORG_STORED
CLIENT_ORG_INFORMATIONAL = {"num_contacts", "contact_names", "emails", "contact_types"}
CLIENT_CONTACT_STORED = set(CONTACT_HEADERS) - {"bank_name", "state"}
CLIENT_CONTACT_INFORMATIONAL = {"bank_name", "state"}


class ParsedWorkbook:
    def __init__(self):
        self.structure_errors = []
        self.groups = []
        self.row_errors = []  # {row_ref, name, field, message, action}
        self.counts = {"organisation_rows": 0, "contact_rows": 0}
        self.sha256 = ""
        self.type_note = ""
        self.format = ""
        self.columns = []  # {sheet, imported, informational, unsupported}

    @property
    def ok(self):
        return not self.structure_errors


def clean(value):
    if value is None:
        return ""

    if isinstance(value, float) and value.is_integer():
        return str(int(value))

    return collapse(value)


def parse_date(value):
    """An observed date as an ISO string (Melbourne midnight), or ''."""
    if value is None or value == "":
        return ""

    if isinstance(value, datetime):
        moment = value if value.tzinfo else value.replace(tzinfo=MELBOURNE)

        return moment.isoformat()

    if isinstance(value, date):
        return datetime(value.year, value.month, value.day, tzinfo=MELBOURNE).isoformat()

    text = clean(value)

    for pattern in ("%Y-%m-%d", "%d/%m/%Y", "%Y-%m-%d %H:%M:%S"):
        try:
            moment = datetime.strptime(text, pattern)

            return moment.replace(tzinfo=MELBOURNE).isoformat()
        except ValueError:
            continue

    return ""


def _read_sheet(worksheet, expected):
    rows = worksheet.iter_rows(values_only=True)

    try:
        header_row = next(rows)
    except StopIteration:
        return None, [], f"The '{worksheet.title}' sheet is empty."

    headers = [clean(value).lower() for value in header_row]
    repeated = sorted({h for h in headers if h and headers.count(h) > 1})

    if repeated:
        return (
            None,
            [],
            f"The '{worksheet.title}' sheet has the column(s) {', '.join(repeated)} more than once, "
            "so it is unclear which values to use. Keep one column of each name.",
        )

    missing = sorted(expected - set(headers))

    if missing:
        return (
            None,
            [],
            f"The '{worksheet.title}' sheet is missing the column(s): "
            + ", ".join(missing)
            + ". Row 1 must contain these headers spelled exactly as shown "
            "(capital letters do not matter; underscores do). The downloadable template has them all.",
        )

    data = []

    for number, values in enumerate(rows, start=2):
        if number > MAX_ROWS:
            return None, [], f"The '{worksheet.title}' sheet has too many rows."

        if not any(clean(v) for v in values):
            continue

        data.append((number, dict(zip(headers, values))))

    return headers, data, ""


def classify_columns(sheet, headers, stored, informational=frozenset()):
    """Which of a sheet's headers are imported, informational only, or unsupported."""
    seen = [h for h in headers if h]

    return {
        "sheet": sheet,
        "imported": [h for h in seen if h in stored],
        "informational": [h for h in seen if h in informational],
        "unsupported": [h for h in seen if h not in stored and h not in informational],
    }


def split_emails(cell):
    """Every email-looking token in a cell (separated by ; , | spaces or new lines), without repeats."""
    seen, tokens = set(), []

    for token in re.split(r"[;,|\s]+", "" if cell is None else str(cell)):
        if token and token.lower() not in seen:
            seen.add(token.lower())
            tokens.append(token)

    return tokens


def split_names(cell):
    """Names separated by ; | or a new line. A comma is never a separator ("Smith, Jo" is one name)."""
    text = "" if cell is None else str(cell)

    return [n for n in (clean_contact_name(part) for part in re.split(r"[;|\r\n]+", text)) if n]


def contacts_from_cells(row_ref, names_cell, emails_cell):
    """
    ``(contacts, notes)`` from the contact_names and emails cells of one row.

    A name is attached to an email only when that is certain: one name and one
    email. With several emails (or several names) nothing is paired by
    position; the contacts are created without names and the combination is
    flagged for review, with the names kept in the message.
    """
    names = split_names(names_cell)
    tokens = split_emails(emails_cell)
    contacts, notes = [], []

    def note(field, message, action, outcome):
        notes.append({"row_ref": row_ref, "field": field, "message": message,
                      "action": action, "outcome": outcome})

    for token in tokens:
        try:
            validate_email(token)
        except ValidationError:
            note("emails", f"'{token}' is not a valid email address, so no contact was created for it.",
                 "Correct the address and import again.", "failed")
            continue

        contacts.append({
            "source_contact_id": "", "row_ref": row_ref, "name": "", "email": token,
            "contact_type": "", "name_needs_verification": None,
            "affiliation_needs_verification": None, "first_observed": "", "last_observed": "",
            "sent_flag": None, "human_reply": None, "status": "",
        })

    if names and not tokens:
        note("contact_names",
             f"Contact name(s) were given ({'; '.join(names)}) without an email address, so no contact was created.",
             "Add the email address(es) and import again.", "failed")
    elif len(names) == 1 and len(tokens) == 1:
        if contacts:
            contacts[0]["name"] = names[0]
    elif names:
        note("contact_names",
             f"{len(names)} name(s) ({'; '.join(names)}) and {len(tokens)} email(s) were given. "
             "Which name belongs to which email cannot be told, so the contacts were added without names.",
             "Open the organisation and enter each contact's name against the right email.", "review")

    return contacts, notes


def _structure_error(result, message):
    result.structure_errors.append(message)

    return result


def _open_sheets(workbook, result):
    """
    Decide the format. Returns ``(format, sheets)`` where sheets maps
    'organisations' (and 'contacts' for the client workbook) to worksheets,
    or ``("", None)`` after recording a structure error.
    """
    names = list(workbook.sheetnames)
    found = "Sheets found: " + (", ".join(f"'{n}'" for n in names) or "none") + "."
    candidates = [n for n in names if n.lower() != "instructions"]

    if SHEET_ORGANISATIONS in names and SHEET_CONTACTS in names:
        return FORMAT_CLIENT, {"organisations": workbook[SHEET_ORGANISATIONS], "contacts": workbook[SHEET_CONTACTS]}

    if SHEET_ORGANISATIONS in names:
        return FORMAT_SIMPLE, {"organisations": workbook[SHEET_ORGANISATIONS]}

    if SHEET_CONTACTS in names:
        result.structure_errors.append(
            f"The workbook has a '{SHEET_CONTACTS}' sheet but no '{SHEET_ORGANISATIONS}' sheet. "
            f"Add the organisations on a sheet named '{SHEET_ORGANISATIONS}'. {found}"
        )

        return "", None

    if len(candidates) == 1:
        return FORMAT_SIMPLE, {"organisations": workbook[candidates[0]]}

    result.structure_errors.append(
        f"The workbook has {'several sheets' if candidates else 'no sheets'} and none is named "
        f"'{SHEET_ORGANISATIONS}'. Put the organisations on one sheet (named '{SHEET_ORGANISATIONS}', "
        f"or keep it the only sheet). A two-sheet client workbook needs '{SHEET_ORGANISATIONS}' and "
        f"'{SHEET_CONTACTS}' spelled exactly like that. {found}"
    )

    return "", None


def parse_workbook(file_bytes, filename=""):
    result = ParsedWorkbook()
    result.sha256 = hashlib.sha256(file_bytes).hexdigest()

    if not filename.lower().endswith(".xlsx"):
        return _structure_error(
            result,
            "Only Microsoft Excel .xlsx files can be imported. A CSV or other spreadsheet "
            "cannot simply be renamed to .xlsx: open it in Excel, arrange it like the template "
            "and use Save As > Excel Workbook (.xlsx).",
        )

    if len(file_bytes) > MAX_BYTES:
        return _structure_error(result, "The file is too large to import.")

    try:
        workbook = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    except Exception:
        return _structure_error(
            result,
            "The file could not be opened as an Excel (.xlsx) workbook. A CSV or other file that "
            "has only been renamed to .xlsx cannot be read: open the original in Excel and save it "
            "as an Excel Workbook (.xlsx) instead.",
        )

    con_rows, con_headers = [], []

    try:
        result.format, sheets = _open_sheets(workbook, result)

        if not result.format:
            return result

        simple = result.format == FORMAT_SIMPLE
        org_sheet = sheets["organisations"]
        org_headers, org_rows, error = _read_sheet(
            org_sheet, set(SIMPLE_HEADERS) if simple else ORGANISATION_COLUMNS
        )

        if error:
            result.structure_errors.append(error)

        if not simple:
            con_headers, con_rows, error = _read_sheet(sheets["contacts"], CONTACT_COLUMNS)

            if error:
                result.structure_errors.append(error)

        if result.structure_errors:
            return result

        if not org_rows:
            result.structure_errors.append(
                f"The '{org_sheet.title}' sheet has no readable rows. If you started from the blank "
                "template, add one row per organisation"
                + ("" if simple else " (and its contacts on the Contacts sheet)")
                + " before uploading; a headers-only file cannot be previewed."
            )

            return result

        sheet_title = org_sheet.title
    finally:
        workbook.close()

    if simple:
        result.columns = [classify_columns(sheet_title, org_headers, SIMPLE_STORED, SIMPLE_INFORMATIONAL)]
    else:
        result.columns = [
            classify_columns(SHEET_ORGANISATIONS, org_headers, CLIENT_ORG_STORED, CLIENT_ORG_INFORMATIONAL),
            classify_columns(SHEET_CONTACTS, con_headers, CLIENT_CONTACT_STORED, CLIENT_CONTACT_INFORMATIONAL),
        ]

    has_type = TYPE_COLUMN in org_headers
    result.counts["organisation_rows"] = len(org_rows)
    result.counts["contact_rows"] = len(con_rows)

    if not has_type:
        result.type_note = (
            "No organisation_type column was supplied, so these records "
            "are imported as Banks (the confirmed community-bank dataset)."
        )

    contacts_by_org = defaultdict(list)

    for number, row in con_rows:
        contacts_by_org[clean(row.get("organisation_id"))].append((number, row))

    seen_ids = set()

    for number, row in org_rows:
        group = _organisation_group(result, row, number, sheet_title, org_headers, has_type, seen_ids)

        if group is None:
            continue

        if simple:
            contacts, problems = contacts_from_cells(group["row_ref"], row.get("contact_names"), row.get("emails"))
            status, reasons = map_claimed_status(group["overall_status_raw"])
            group["contact_names_in_file"] = clean(row.get("contact_names"))
            group["emails_in_file"] = clean(row.get("emails"))
        else:
            contacts, problems = _client_contacts(contacts_by_org.get(group["source_org_id"], []))
            status, reasons = map_organisation_status(contacts)

        group.update(
            format=result.format, contacts=contacts, contact_problems=problems,
            evidence_status=status, evidence_reasons=reasons,
        )
        result.groups.append(group)

    known_ids = {clean(r.get("organisation_id")) for _n, r in org_rows}

    for source_id, rows in contacts_by_org.items():
        if source_id not in known_ids:
            for number, row in rows:
                result.row_errors.append(
                    {
                        "row_ref": f"Contacts row {number}",
                        "source_id": source_id,
                        "name": clean(row.get("bank_name")),
                        "field": "organisation_id",
                        "message": (
                            "The contact refers to an organisation that is not in the Organisations sheet."
                            if source_id else "The contact has no organisation ID."
                        ),
                        "action": "Add the organisation or correct the ID.",
                    }
                )

    return result


def _organisation_group(result, row, number, sheet_title, org_headers, has_type, seen_ids):
    """One organisation's validated fields, or None after recording why the row was rejected."""
    source_id = clean(row.get("organisation_id"))
    name = clean(row.get("bank_name"))
    row_ref = f"{sheet_title} row {number}"

    def reject(field, message, action, outcome="failed"):
        result.row_errors.append(
            {
                "row_ref": row_ref,
                "source_id": source_id,
                "name": name,
                "field": field,
                "message": message,
                "action": action,
                "outcome": outcome,
            }
        )

    if not source_id:
        reject("organisation_id", "The row has no organisation ID.",
               "Add the organisation ID and import again.")
        return None

    if source_id in seen_ids:
        reject("organisation_id",
               f"The organisation ID '{source_id}' appears more than once.",
               "Remove the repeated row.")
        return None

    seen_ids.add(source_id)

    if not name:
        reject("bank_name", "The organisation name is missing.",
               "Add the organisation name.")
        return None

    state = norm_state(row.get("state"))

    if not state:
        reject("state", "The state or territory is missing or not recognised.",
               "Enter a valid Australian state or territory.")
        return None

    raw_type = clean(row.get(TYPE_COLUMN)).lower() if has_type else "bank"
    org_type = SUPPORTED_TYPES.get(raw_type)

    if org_type is None:
        reject(TYPE_COLUMN,
               f"The organisation type '{raw_type or '(blank)'}' is not supported, so the record was held for review and not imported.",
               "Correct the type to Bank, Branch or Club and import the file again.",
               outcome="review")
        return None

    optional = {}

    for column, attribute in OPTIONAL_COLUMNS.items():
        if column in org_headers:
            value = clean(row.get(column))

            if value:
                optional[attribute] = value

    return {
        "source_org_id": source_id,
        "row_ref": row_ref,
        "name": name,
        "state": state,
        "org_type": org_type,
        "optional": optional,
        "overall_status_raw": clean(row.get("overall_status")),
        "needs_verification": parse_bool(row.get("needs_verification_flag")),
        "any_human_reply": parse_bool(row.get("any_human_reply_received")),
        "all_contacts_replied": parse_bool(row.get("all_contacts_replied")),
        "first_observed": parse_date(row.get("first_observed_date")),
        "last_observed": parse_date(row.get("last_observed_date")),
    }


def _client_contacts(rows):
    """Contacts of the two-sheet workbook for one organisation: ``(contacts, problems)``."""
    contacts, problems, seen_contact_ids = [], [], set()

    for c_number, c_row in rows:
        contact_id = clean(c_row.get("contact_id"))
        email = clean(c_row.get("email"))
        c_ref = f"Contacts row {c_number}"

        if not contact_id or contact_id in seen_contact_ids:
            problems.append({"row_ref": c_ref, "field": "contact_id",
                             "message": "The contact ID is missing or repeated."})
            continue

        seen_contact_ids.add(contact_id)

        if not email:
            problems.append({"row_ref": c_ref, "field": "email",
                             "message": "The contact has no email address."})
            continue

        try:
            validate_email(email)
        except ValidationError:
            problems.append({"row_ref": c_ref, "field": "email",
                             "message": "The contact email address is not valid."})
            continue

        contacts.append(
            {
                "source_contact_id": contact_id,
                "row_ref": c_ref,
                "name": clean_contact_name(c_row.get("contact_name")),
                "email": email,
                "contact_type": clean(c_row.get("contact_type")),
                "name_needs_verification": parse_bool(c_row.get("name_needs_verification")),
                "affiliation_needs_verification": parse_bool(
                    c_row.get("bank_affiliation_needs_verification")
                ),
                "first_observed": parse_date(c_row.get("first_observed_date")),
                "last_observed": parse_date(c_row.get("last_observed_date")),
                "sent_flag": parse_bool(c_row.get("sent_to_by_srfu")),
                "human_reply": parse_bool(c_row.get("human_reply_received")),
                "status": clean(c_row.get("contact_status")),
            }
        )

    return contacts, problems
