"""
The downloadable import template, a synthetic example and the on-page
guidance. Everything here is derived from the parser's own rules
(``workbook.py``, ``evidence.py``, ``normalise.py``) so the guidance cannot
drift from what the import actually accepts.

No client data, credentials or real addresses are used. Example rows use
invented names and ``example.test`` addresses.
"""

import io

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .evidence import EXPECTED_FLAGS, KNOWN, SUCCESSFUL
from .normalise import PLACEHOLDER_NAMES, STATE_NAMES
from .workbook import (
    CONTACT_HEADERS,
    MAX_BYTES,
    MAX_ROWS,
    OPTIONAL_COLUMNS,
    ORGANISATION_HEADERS,
    SHEET_CONTACTS,
    SHEET_ORGANISATIONS,
    TYPE_COLUMN,
)

TEMPLATE_FILENAME = "SRFU-import-template.xlsx"
EXAMPLE_FILENAME = "SRFU-import-EXAMPLE-synthetic-do-not-import.xlsx"
SHEET_INSTRUCTIONS = "Instructions"

# Optional organisation columns in the order they are offered in the template.
OPTIONAL_ORGANISATION_HEADERS = (TYPE_COLUMN,) + tuple(OPTIONAL_COLUMNS)

# What each header is for. Columns marked "informational" must exist but their
# values do not decide matching (checked against the parser).
ORGANISATION_NOTES = {
    "organisation_id": "Required value. Your own unique identifier for the organisation. Contacts link to it. Must not repeat.",
    "bank_name": "Required value. The organisation's name (for a Branch or Club, its own name).",
    "state": "Required value. Australian state or territory: ACT, NSW, NT, QLD, SA, TAS, VIC or WA (full names such as Victoria also work).",
    "num_contacts": "Header required; value informational only (not used).",
    "contact_names": "Header required; value informational only (contacts are read from the Contacts sheet).",
    "emails": "Header required; value informational only (contacts are read from the Contacts sheet).",
    "contact_types": "Header required; value informational only.",
    "any_human_reply_received": "Header required. TRUE or FALSE. Kept as reply evidence; it never sets Interested.",
    "all_contacts_replied": "Header required. TRUE or FALSE. Kept as evidence only.",
    "overall_status": "Header required. Kept as supplied for reference. It is NOT used to set the contact status (see contact_status on the Contacts sheet).",
    "first_observed_date": "Header required. Date first observed (blank allowed).",
    "last_observed_date": "Header required. Date last observed (blank allowed).",
    "needs_verification_flag": "Header required. TRUE if the record still needs checking (blank allowed).",
}
OPTIONAL_NOTES = {
    TYPE_COLUMN: "Bank, Branch or Club. If this column is absent every row is imported as a Bank. If present it must be filled in: a blank or other value holds the row for review.",
    "parent_bank_name": "Branches: the parent bank's name. It must match a bank that already exists in the system, otherwise the branch is reported as incomplete.",
    "parent_bank_state": "Branches: the parent bank's state (defaults to the row's state).",
    "region": "Optional region.",
    "suburb": "Optional suburb. Used for Club and Branch duplicate checks.",
    "postcode": "Optional four-digit postcode. Used for Club and Branch duplicate checks.",
    "address": "Optional street address.",
    "website_url": "Optional public website.",
    "public_phone": "Optional public phone number.",
    "public_email": "Optional public organisation email (otherwise the first contact's email is used).",
    "club_type": "Clubs: optional sport or club type.",
}
CONTACT_NOTES = {
    "contact_id": "Required value. Unique per contact; not repeated within an organisation.",
    "organisation_id": "Required value. Must equal an organisation_id on the Organisations sheet; this is how a contact is linked to its organisation.",
    "bank_name": "Header required; value informational only (the link is organisation_id).",
    "state": "Header required; value informational only.",
    "contact_name": "Optional. Leave blank when unknown. '(unspecified)' is treated as blank.",
    "email": "Required value for the contact, in a valid email format. A contact without a valid email is reported and skipped; the organisation is still imported.",
    "contact_type": "Optional, for example shared_mailbox or named.",
    "name_needs_verification": "Header required. TRUE or FALSE (blank allowed).",
    "bank_affiliation_needs_verification": "Header required. TRUE or FALSE (blank allowed).",
    "first_observed_date": "Header required. Date (blank allowed).",
    "last_observed_date": "Header required. Date (blank allowed).",
    "sent_to_by_srfu": "Header required. TRUE or FALSE. Must agree with contact_status (see the table below).",
    "human_reply_received": "Header required. TRUE or FALSE. Must agree with contact_status.",
    "contact_status": "Outreach evidence for this contact: one of the values below, or blank for no history. It decides whether the organisation is Contacted.",
}

STATUS_MEANING = {
    "sent_and_human_reply_received": "Contacted (a confirmed outgoing email)",
    "sent_no_human_reply_found": "Contacted (a confirmed outgoing email)",
    "delivery_failure_recorded": "Not Yet Contacted (delivery did not succeed)",
    "human_inbound_only": "Not Yet Contacted (nothing was sent by SRFU)",
    "copied_or_automated_only": "Not Yet Contacted (not enough evidence of outreach)",
    "no_outreach_history": "Not Yet Contacted",
}


def _flag(value):
    return "either" if value is None else ("TRUE" if value else "FALSE")


def guidance():
    """The template context for the Import Data page (all derived from the rules)."""
    statuses = [
        {
            "value": status,
            "result": STATUS_MEANING[status],
            "sent": _flag(EXPECTED_FLAGS[status]["sent"]),
            "reply": _flag(EXPECTED_FLAGS[status]["reply"]),
        }
        for status in KNOWN
    ]

    return {
        "sheet_organisations": SHEET_ORGANISATIONS,
        "sheet_contacts": SHEET_CONTACTS,
        "organisation_columns": [(h, ORGANISATION_NOTES[h]) for h in ORGANISATION_HEADERS],
        "optional_columns": [(h, OPTIONAL_NOTES[h]) for h in OPTIONAL_ORGANISATION_HEADERS],
        "contact_columns": [(h, CONTACT_NOTES[h]) for h in CONTACT_HEADERS],
        "statuses": statuses,
        "states": sorted({code for code in STATE_NAMES.values()}),
        "placeholder_names": sorted(PLACEHOLDER_NAMES),
        "successful": SUCCESSFUL,
        "max_mb": MAX_BYTES // (1024 * 1024),
        "max_rows": MAX_ROWS,
        "template_filename": TEMPLATE_FILENAME,
        "example_filename": EXAMPLE_FILENAME,
    }


# ---------------------------------------------------------------- workbooks
def _header(sheet, headers):
    sheet.append(list(headers))
    bold = Font(bold=True)
    fill = PatternFill("solid", fgColor="DCE6F1")

    for index, cell in enumerate(sheet[1], start=1):
        cell.font = bold
        cell.fill = fill
        sheet.column_dimensions[get_column_letter(index)].width = max(len(str(cell.value)) + 4, 14)

    sheet.freeze_panes = "A2"


def _instructions(sheet, example):
    lines = [
        ("SRFU data import " + ("EXAMPLE (synthetic data - do not import)" if example else "template"), True),
        ("", False),
    ]

    if example:
        lines += [
            ("Every name, email and identifier here is invented (example.test addresses). It exists only to show the layout.", False),
            ("Do not import this file into the real system. Download the blank template and enter real records there.", False),
        ]
    else:
        lines += [
            ("This workbook contains only headers. Add your records, save, then upload it on the Import Data page.", False),
            ("A headers-only file cannot be previewed: the Organisations sheet needs at least one data row.", False),
        ]

    lines += [
        ("", False),
        (f"Keep the sheets named exactly '{SHEET_ORGANISATIONS}' and '{SHEET_CONTACTS}'. This Instructions sheet is ignored.", False),
        ("Keep row 1 (the headers) exactly as supplied; do not rename, translate or add spaces to the headers.", False),
        ("One row per organisation on 'Organisations'; one row per contact on 'Contacts'.", False),
        ("Link a contact to its organisation by giving the contact the same organisation_id.", False),
        ("Several contacts may share one organisation_id; the organisation is created once.", False),
        ("Leave unknown optional cells empty. Do not type placeholders such as 'n/a'.", False),
        ("Formulas must have been calculated and saved by Excel; type plain values.", False),
        ("The Import Data page lists every column, the accepted values and what happens to bad rows.", False),
    ]

    for text, bold in lines:
        sheet.append([text])

        if bold:
            sheet.cell(row=sheet.max_row, column=1).font = Font(bold=True, size=13)

    sheet.column_dimensions["A"].width = 110


def _workbook(organisations, contacts, example):
    wb = Workbook()
    wb.remove(wb.active)
    org_sheet = wb.create_sheet(SHEET_ORGANISATIONS)
    _header(org_sheet, ORGANISATION_HEADERS + OPTIONAL_ORGANISATION_HEADERS)

    for row in organisations:
        org_sheet.append(row)

    con_sheet = wb.create_sheet(SHEET_CONTACTS)
    _header(con_sheet, CONTACT_HEADERS)

    for row in contacts:
        con_sheet.append(row)

    _instructions(wb.create_sheet(SHEET_INSTRUCTIONS), example)
    buffer = io.BytesIO()
    wb.save(buffer)

    return buffer.getvalue()


def build_template():
    """A blank workbook: headers only, no records."""
    return _workbook([], [], example=False)


def build_example():
    """A small, clearly synthetic workbook that previews and imports cleanly."""
    d1, d2 = "2025-01-10", "2025-02-10"
    org_pad = len(OPTIONAL_ORGANISATION_HEADERS)

    def org(oid, name, state, kind, status, reply, **optional):
        row = [oid, name, state, "", "", "", "", reply, "FALSE", status, d1, d2, "TRUE"]
        extras = {TYPE_COLUMN: kind, **optional}

        return row + [extras.get(h, "") for h in OPTIONAL_ORGANISATION_HEADERS][:org_pad]

    organisations = [
        org("EXAMPLE_1", "EXAMPLE ONLY Alpha Community Bank", "VIC", "Bank", "Contacted", "TRUE"),
        org("EXAMPLE_2", "EXAMPLE ONLY Bravo Community Bank", "NSW", "Bank", "Not Yet Contacted", "FALSE"),
        org("EXAMPLE_3", "EXAMPLE ONLY Charlie Netball Club", "VIC", "Club", "Not Yet Contacted", "FALSE",
            suburb="Exampleton", postcode="3999", club_type="Netball"),
    ]

    def contact(cid, oid, name, email, ctype, sent, reply, status, state):
        return [cid, oid, "", state, name, email, ctype, "FALSE", "FALSE", d1, d2, sent, reply, status]

    contacts = [
        contact("EXAMPLE_C1", "EXAMPLE_1", "Pat Example", "pat@example.test", "named",
                "TRUE", "TRUE", "sent_and_human_reply_received", "VIC"),
        contact("EXAMPLE_C2", "EXAMPLE_1", "", "info@example.test", "shared_mailbox",
                "TRUE", "FALSE", "sent_no_human_reply_found", "VIC"),
        contact("EXAMPLE_C3", "EXAMPLE_2", "", "bravo@example.test", "shared_mailbox",
                "FALSE", "FALSE", "no_outreach_history", "NSW"),
    ]

    return _workbook(organisations, contacts, example=True)
